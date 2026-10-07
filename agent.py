import asyncio
import base64
import io
import re
import time
import xml.etree.ElementTree as ET
import json
import logging
import os
import smtplib
import ssl
import sys
import urllib.parse
import webbrowser
from datetime import datetime
from html.parser import HTMLParser
from email.message import EmailMessage
from pathlib import Path

import httpx
from ddgs import DDGS
from dotenv import load_dotenv
from livekit.agents import (
    Agent,
    AgentServer,
    AgentSession,
    JobContext,
    RunContext,
    cli,
    function_tool,
    room_io,
)
from livekit.plugins import google

GUI_ERROR = ""
try:
    import pyautogui

    if sys.platform == "win32":
        try:
            import ctypes

            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:
            pass
    pyautogui.FAILSAFE = True  # slam the mouse into a screen corner to abort
    pyautogui.PAUSE = 0.15
    HAVE_GUI = True
except Exception as e:
    HAVE_GUI = False
    GUI_ERROR = repr(e)

try:
    from docx import Document
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

load_dotenv(".env.local")
logger = logging.getLogger("cortana")

# Add people here so you can just say "email mom". Names must be lowercase.
CONTACTS = {
    "example": "example@gmail.com",
}

# Websites she can open by name. Add your own.
SITES = {
    "google": "https://www.google.com",
    "youtube": "https://www.youtube.com",
    "gmail": "https://mail.google.com",
    "maps": "https://maps.google.com",
    "netflix": "https://www.netflix.com",
    "reddit": "https://www.reddit.com",
    "github": "https://github.com",
    "amazon": "https://www.amazon.com",
    "spotify": "https://open.spotify.com",
}

# Sites where she can also run a search.
SEARCH_URLS = {
    "google": "https://www.google.com/search?q=",
    "youtube": "https://www.youtube.com/results?search_query=",
    "amazon": "https://www.amazon.com/s?k=",
    "maps": "https://www.google.com/maps/search/",
}

# Apps she can open (Windows). She can ONLY open apps on this list. Add your own.
APPS = {
    "notepad": "notepad.exe",
    "calculator": "calc.exe",
    "file explorer": "explorer.exe",
    "task manager": "taskmgr.exe",
    "paint": "mspaint.exe",
    "settings": "ms-settings:",
    "chrome": "chrome.exe",
    "edge": "msedge.exe",
    "word": "winword.exe",
    "excel": "excel.exe",
    "spotify": "spotify:",
    "discord": "discord:",
    "steam": "steam://open/main",
}

INSTRUCTIONS = """You are Cortana, a sharp, calm, quick-witted AI companion in the style of a futuristic tactical assistant.
Personality: confident, composed, dryly funny, warm underneath, and loyal to the user. You tease lightly but always help.
Speak in crisp, natural sentences, one to three at a time. No lists, symbols, or markdown.
Use the web search tool for anything current, like news, weather, prices, or scores. Say a short line like "Pulling that up" first, then answer in your own words.
To send an email, first read back the recipient, subject, and message, then ask "Should I send it?" Only send after a clear yes. Never guess an email address.
You have long-term memory sorted into categories: fact (name, birthday, job), preference (likes, dislikes, how they want to be treated), person (someone in their life and who they are), routine (a habit or regular schedule), and project (something ongoing they are working toward). When the user shares something lasting, or asks you to remember something, save it with the remember tool using the best category. Keep each memory to one short plain sentence. Never save passwords, keys, or payment details. If asked to forget something, use the forget tool. If something you remember turns out to be wrong or outdated, use update_memory rather than saving a duplicate. Use what you remember naturally in conversation and do not list it unless asked. You may occasionally and gently reference a memory if it genuinely fits what is being discussed, but never force it in.
Use get_current_time before working out reminder times. Set reminders and timers with set_reminder, and keep to-do and shopping lists with the list tools.
For a daily briefing, use daily_briefing. Use the user's remembered home city for weather, or ask for it once and remember it.
You can control volume and music playback with media_control.
For big questions that need real research, use the research tool, and tell the user first that it will take about a minute. For writing jobs like essays, letters, plans, or summaries, use write_document. Afterward give a short spoken summary. Never read a whole document aloud.
Computer control: only when the user asks, call enable_control. Then you can use look_at_screen, click_on, type_text, press_keys, and scroll. Work in small steps and check the screen after important steps. Before anything that sends, deletes, buys, posts, or closes something, say exactly what you are about to do, ask "Should I do it?", and continue only after a clear yes. Never type passwords or payment details; ask the user to do that themselves. If the user says stop, call stop_control right away.
You can open websites and apps on the user's computer with your tools. Just do it when asked, then confirm in a few words.
You can see through the user's camera when it is on. If they ask what you see and the camera is off, ask them to turn it on.
Never claim to be the actual character from any game. You are an original assistant with a similar attitude."""


MEMORY_FILE = Path(__file__).parent / "memory.json"


def load_memories() -> list:
    """Each memory is {text, category, created, last_used, use_count}. Old plain-string
    files are upgraded automatically the first time they are loaded."""
    try:
        items = json.loads(MEMORY_FILE.read_text(encoding="utf-8"))
    except Exception:
        return []
    changed = False
    for i, m in enumerate(items):
        if isinstance(m, str):
            items[i] = {"text": m, "category": "fact", "created": time.time(), "last_used": time.time(), "use_count": 0}
            changed = True
    if changed:
        save_memories(items)
    return items


def save_memories(items: list) -> None:
    MEMORY_FILE.write_text(json.dumps(items, indent=2), encoding="utf-8")


def memory_texts(items: list) -> list:
    return [m["text"] for m in items]


def log_event(kind: str, detail: str) -> None:
    """Append a short line to the day's activity journal, used to notice patterns."""
    try:
        items = load_json("journal.json", [])
        items.append({"t": time.time(), "kind": kind, "detail": detail[:200]})
        save_json("journal.json", items[-500:])
    except Exception:
        logger.exception("journal write failed")


DATA_DIR = Path(__file__).parent


def load_json(name, default):
    try:
        return json.loads((DATA_DIR / name).read_text(encoding="utf-8"))
    except Exception:
        return default


def save_json(name, data) -> None:
    (DATA_DIR / name).write_text(json.dumps(data, indent=2), encoding="utf-8")


WEATHER_WORDS = {
    0: "clear", 1: "mostly clear", 2: "partly cloudy", 3: "cloudy", 45: "foggy", 48: "foggy",
    51: "light drizzle", 53: "drizzle", 55: "heavy drizzle", 61: "light rain", 63: "rain", 65: "heavy rain",
    71: "light snow", 73: "snow", 75: "heavy snow", 80: "rain showers", 81: "rain showers",
    82: "heavy showers", 95: "thunderstorms", 96: "thunderstorms with hail", 99: "thunderstorms with hail",
}

MEDIA_KEYS = {"mute": 0xAD, "next": 0xB0, "previous": 0xB1, "stop": 0xB2, "play_pause": 0xB3}


DOCS_DIR = DATA_DIR / "documents"


async def gemini_text(prompt: str, system: str = "", image_jpeg: bytes = None) -> str:
    """Ask a Gemini text model to write something. Tries several model names."""
    key = os.environ["GOOGLE_API_KEY"]
    models = [m for m in [os.getenv("GEMINI_TEXT_MODEL"), "gemini-3.1-flash-lite",
                          "gemini-3-flash-preview", "gemini-2.5-flash"] if m]
    parts = [{"text": prompt}]
    if image_jpeg:
        parts.insert(0, {"inline_data": {"mime_type": "image/jpeg", "data": base64.b64encode(image_jpeg).decode()}})
    body = {"contents": [{"parts": parts}]}
    if system:
        body["systemInstruction"] = {"parts": [{"text": system}]}
    last = ""
    async with httpx.AsyncClient(timeout=120) as client:
        for m in models:
            r = await client.post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{m}:generateContent",
                headers={"x-goog-api-key": key}, json=body)
            if r.status_code == 200:
                try:
                    return "".join(p.get("text", "") for p in r.json()["candidates"][0]["content"]["parts"])
                except Exception:
                    last = f"{m}: empty answer"
                    continue
            last = f"{m}: {r.status_code} {r.text[:150]}"
    raise RuntimeError(last)


class _PageText(HTMLParser):
    SKIP = ("script", "style", "nav", "footer", "header", "aside", "noscript", "svg", "form")

    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip and data.strip():
            self.parts.append(data.strip())


async def fetch_page(client, url: str) -> str:
    try:
        r = await client.get(url, headers={"User-Agent": "Mozilla/5.0"}, follow_redirects=True)
        if "text/html" not in r.headers.get("content-type", ""):
            return ""
        p = _PageText()
        p.feed(r.text)
        return " ".join(p.parts)[:3500]
    except Exception:
        return ""


def _add_runs(paragraph, text: str) -> None:
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            paragraph.add_run(chunk.replace("`", "")).bold = bool(i % 2)


def save_document(title: str, markdown: str) -> Path:
    DOCS_DIR.mkdir(exist_ok=True)
    safe = re.sub(r"[^\w\- ]", "", title).strip()[:50] or "Document"
    stamp = datetime.now().strftime("%Y-%m-%d %H%M")
    if not HAVE_DOCX:
        path = DOCS_DIR / f"{safe} {stamp}.md"
        path.write_text(markdown, encoding="utf-8")
        return path
    doc = Document()
    doc.add_heading(title, 0)
    for line in markdown.splitlines():
        t = line.rstrip()
        if not t or t.startswith("# "):
            continue
        if t.startswith("### "):
            doc.add_heading(t[4:], 3)
        elif t.startswith("## "):
            doc.add_heading(t[3:], 1)
        elif t.startswith(("- ", "* ")):
            _add_runs(doc.add_paragraph(style="List Bullet"), t[2:])
        else:
            _add_runs(doc.add_paragraph(), t)
    path = DOCS_DIR / f"{safe} {stamp}.docx"
    doc.save(path)
    return path


def section(markdown: str, name: str) -> str:
    m = re.search(rf"^##\s*{name}\s*\n(.*?)(?=^##\s|\Z)", markdown, re.S | re.M | re.I)
    return m.group(1).strip() if m else ""


RISKY_WORDS = (
    "send", "submit", "delete", "remove", "buy", "purchase", "pay", "order", "checkout", "confirm",
    "close", "quit", "exit", "uninstall", "install", "post", "publish", "sign out", "log out", "logout",
    "shut down", "restart", "format", "erase", "empty", "transfer",
)


def click_is_risky(description: str) -> bool:
    d = description.lower()
    return any(w in d for w in RISKY_WORDS)


def keys_are_risky(parts) -> bool:
    p = set(parts)
    return "delete" in p or "win" in p or ("alt" in p and "f4" in p) or ("ctrl" in p and "w" in p)


def parse_point(text: str):
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        return None
    try:
        d = json.loads(m.group(0))
        if not d.get("found"):
            return None
        x, y = float(d["x"]), float(d["y"])
    except Exception:
        return None
    if not (0 <= x <= 1000 and 0 <= y <= 1000):
        return None
    return x / 1000, y / 1000


def screenshot_jpeg() -> bytes:
    img = pyautogui.screenshot()
    w = img.size[0]
    if w > 1600:
        img = img.resize((1600, int(img.size[1] * 1600 / w)))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=80)
    return buf.getvalue()


class Cortana(Agent):
    def __init__(self, room=None) -> None:
        memories = load_memories()
        extra = ""
        if memories:
            by_cat = {}
            for m in memories:
                by_cat.setdefault(m.get("category", "fact"), []).append(m["text"])
            blocks = []
            for cat, label in (("fact", "Facts about the user"), ("preference", "Preferences"),
                               ("person", "People in their life"), ("routine", "Routines and habits"),
                               ("project", "Ongoing projects or goals")):
                if by_cat.get(cat):
                    blocks.append(label + ":\n" + "\n".join(f"- {t}" for t in by_cat[cat]))
            extra = "\n\nWhat you remember about the user:\n" + "\n\n".join(blocks)
        super().__init__(instructions=INSTRUCTIONS + extra)
        self._tasks = set()
        self._room = room
        self._armed = False
        self._pending = None

    async def on_enter(self):
        now = time.time()
        reminders = load_json("reminders.json", [])
        due = [r for r in reminders if r["due"] <= now]
        pending = [r for r in reminders if r["due"] > now]
        if due:
            save_json("reminders.json", pending)
        for r in pending:
            self._schedule(r)
        text = "Greet the user briefly and offer your help."
        if due:
            text += " Also tell them about these reminders that came due while you were offline: " + "; ".join(r["message"] for r in due)
        note = self._observation()
        if note:
            text += " " + note
        await self.session.generate_reply(instructions=text)

    def _observation(self) -> str:
        """Look for a pattern worth mentioning: a long gap since last talking, or an old
        open project that has not come up in a while. Keeps it to at most one gentle note."""
        try:
            journal = load_json("journal.json", [])
            last_session = max((e["t"] for e in journal if e["kind"] == "session"), default=None)
            log_event("session", "opened")
            gap_note = ""
            if last_session and time.time() - last_session > 60 * 60 * 24 * 4:
                days = int((time.time() - last_session) / 86400)
                gap_note = f"It has been about {days} days since you last talked. You may want to say something like a friendly welcome back, but keep it brief and do not make a big deal of it."
            projects = [m for m in load_memories() if m.get("category") == "project"]
            stale = [m for m in projects if time.time() - m.get("last_used", 0) > 60 * 60 * 24 * 7]
            project_note = ""
            if stale and not gap_note:
                project_note = f"You could casually ask about this if it fits naturally, but only once and only if it is a good moment: {stale[0]['text']}"
            return " ".join(x for x in (gap_note, project_note) if x)
        except Exception:
            logger.exception("observation failed")
            return ""

    def _schedule(self, r):
        task = asyncio.create_task(self._fire(r))
        self._tasks.add(task)
        task.add_done_callback(self._tasks.discard)

    async def _fire(self, r):
        await asyncio.sleep(max(0, r["due"] - time.time()))
        items = load_json("reminders.json", [])
        if not any(x["id"] == r["id"] for x in items):
            return
        save_json("reminders.json", [x for x in items if x["id"] != r["id"]])
        try:
            await self.session.generate_reply(
                instructions=f"Speak up now and tell the user this reminder: {r['message']}"
            )
        except Exception:
            logger.exception("reminder failed")

    @function_tool
    async def search_web(self, context: RunContext, query: str) -> str:
        """Search the web for current information.

        Args:
            query: A short, specific search query.
        """
        def _search():
            with DDGS() as ddgs:
                return list(ddgs.text(query, max_results=5))

        try:
            results = await asyncio.to_thread(_search)
        except Exception as e:
            return f"Search failed: {e}"
        if not results:
            return "No results found."
        return "\n\n".join(f"{r.get('title', '')}: {r.get('body', '')}" for r in results)


    @function_tool
    async def send_email(self, context: RunContext, to: str, subject: str, body: str) -> str:
        """Send an email from the user's account. Only call this after the user has confirmed
        the recipient, subject, and message out loud.

        Args:
            to: A contact name or a full email address.
            subject: The email subject line.
            body: The full email text.
        """
        address = CONTACTS.get(to.strip().lower(), to.strip())
        if "@" not in address:
            return "There is no email address for that name. Ask the user for the address."
        user = os.getenv("EMAIL_ADDRESS")
        password = os.getenv("EMAIL_APP_PASSWORD")
        if not user or not password:
            return "Email is not set up yet. Tell the user to run the email setup file."

        def _send():
            msg = EmailMessage()
            msg["From"] = user
            msg["To"] = address
            msg["Subject"] = subject
            msg.set_content(body)
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as smtp:
                smtp.login(user, password)
                smtp.send_message(msg)

        try:
            await asyncio.to_thread(_send)
        except Exception as e:
            return f"Sending failed: {e}"
        return f"Email sent to {address}."

    @function_tool
    async def open_website(self, context: RunContext, site: str, query: str = "") -> str:
        """Open a website in the user's browser, optionally searching it.

        Args:
            site: A site name like youtube or gmail, or a web address like example.com.
            query: Optional search words, for example "lofi music" to search on youtube.
        """
        key = site.strip().lower()
        if query and key in SEARCH_URLS:
            url = SEARCH_URLS[key] + urllib.parse.quote_plus(query)
        elif key in SITES:
            url = SITES[key]
        elif "." in key and " " not in key:
            url = key if key.startswith(("http://", "https://")) else "https://" + key
        else:
            return "I don't know that site. Ask the user for the web address."
        if not url.startswith(("http://", "https://")):
            return "Only normal web addresses can be opened."
        await asyncio.to_thread(webbrowser.open, url)
        return f"Opened {site}."

    @function_tool
    async def open_app(self, context: RunContext, app: str) -> str:
        """Open an app on the user's computer from the allowed list.

        Args:
            app: The app name, for example notepad, calculator, or spotify.
        """
        target = APPS.get(app.strip().lower())
        if target is None:
            return "That app is not on my allowed list. The user can add it to the APPS list."
        if sys.platform != "win32":
            return "Opening apps only works on Windows."
        try:
            os.startfile(target)
        except Exception as e:
            return f"Could not open {app}: {e}"
        return f"Opened {app}."

    @function_tool
    async def remember(self, context: RunContext, fact: str, category: str = "fact") -> str:
        """Save a lasting fact about the user to long-term memory.

        Args:
            fact: One short plain sentence, for example "The user's name is Ethan."
            category: One of fact, preference, person, routine, project.
        """
        fact = fact.strip()[:300]
        if not fact:
            return "Nothing to save."
        category = category.strip().lower()
        if category not in ("fact", "preference", "person", "routine", "project"):
            category = "fact"
        items = load_memories()
        if fact.lower() in [m["text"].lower() for m in items]:
            return "Already remembered."
        items.append({"text": fact, "category": category, "created": time.time(),
                      "last_used": time.time(), "use_count": 0})
        save_memories(items[-300:])
        log_event("remember", fact)
        return "Saved."

    @function_tool
    async def update_memory(self, context: RunContext, old_topic: str, new_fact: str, category: str = "fact") -> str:
        """Replace an outdated memory with a corrected one.

        Args:
            old_topic: A word or phrase that appears in the memory to replace.
            new_fact: The corrected one-sentence fact.
            category: One of fact, preference, person, routine, project.
        """
        items = load_memories()
        topic = old_topic.strip().lower()
        kept = [m for m in items if topic not in m["text"].lower()]
        removed = len(items) - len(kept)
        category = category.strip().lower()
        if category not in ("fact", "preference", "person", "routine", "project"):
            category = "fact"
        kept.append({"text": new_fact.strip()[:300], "category": category, "created": time.time(),
                     "last_used": time.time(), "use_count": 0})
        save_memories(kept[-300:])
        return f"Updated. Replaced {removed} old item(s)."

    @function_tool
    async def forget(self, context: RunContext, topic: str) -> str:
        """Delete saved memories that mention a topic.

        Args:
            topic: A word or phrase that appears in the memory to delete.
        """
        topic = topic.strip().lower()
        if not topic:
            return "Nothing to forget."
        items = load_memories()
        kept = [m for m in items if topic not in m["text"].lower()]
        save_memories(kept)
        removed = len(items) - len(kept)
        return f"Forgot {removed} item(s)." if removed else "I had nothing saved about that."

    @function_tool
    async def recall(self, context: RunContext, topic: str = "") -> str:
        """Search long-term memory for anything related to a topic. Use this if you are not
        sure whether something has already been remembered, instead of guessing.

        Args:
            topic: A word or phrase to search for. Leave empty to get everything.
        """
        items = load_memories()
        topic = topic.strip().lower()
        matches = items if not topic else [m for m in items if topic in m["text"].lower()]
        if not matches:
            return "Nothing found."
        now = time.time()
        for m in matches:
            m["last_used"] = now
            m["use_count"] = m.get("use_count", 0) + 1
        save_memories(items)
        return "\n".join(f"[{m.get('category', 'fact')}] {m['text']}" for m in matches)

    @function_tool
    async def get_current_time(self, context: RunContext) -> str:
        """Get the current date and time on the user's computer."""
        return datetime.now().strftime("%A, %B %d, %Y, %I:%M %p")

    @function_tool
    async def set_reminder(self, context: RunContext, minutes: float, message: str) -> str:
        """Set a reminder or timer that speaks up after some minutes.

        Args:
            minutes: How many minutes from now. Work this out with get_current_time if needed.
            message: What to remind the user about.
        """
        if minutes <= 0 or minutes > 60 * 24 * 30:
            return "That time is not valid."
        r = {"id": str(int(time.time() * 1000)), "due": time.time() + minutes * 60, "message": message.strip()[:200]}
        items = load_json("reminders.json", [])
        items.append(r)
        save_json("reminders.json", items)
        self._schedule(r)
        return f"Reminder set for {minutes:g} minutes from now."

    @function_tool
    async def list_reminders(self, context: RunContext) -> str:
        """List the user's upcoming reminders."""
        items = load_json("reminders.json", [])
        if not items:
            return "No reminders."
        now = time.time()
        return "\n".join(f"{r['message']} in {max(0, round((r['due'] - now) / 60))} minutes" for r in items)

    @function_tool
    async def cancel_reminder(self, context: RunContext, topic: str) -> str:
        """Cancel reminders whose text mentions a topic.

        Args:
            topic: A word from the reminder to cancel.
        """
        items = load_json("reminders.json", [])
        kept = [r for r in items if topic.strip().lower() not in r["message"].lower()]
        save_json("reminders.json", kept)
        return f"Cancelled {len(items) - len(kept)} reminder(s)."

    @function_tool
    async def add_to_list(self, context: RunContext, list_name: str, item: str) -> str:
        """Add an item to a named list such as todo or shopping.

        Args:
            list_name: The list name, for example todo or shopping.
            item: The item to add.
        """
        lists = load_json("lists.json", {})
        key = list_name.strip().lower()
        lists.setdefault(key, []).append(item.strip()[:200])
        save_json("lists.json", lists)
        return f"Added to {key}."

    @function_tool
    async def read_list(self, context: RunContext, list_name: str) -> str:
        """Read all items on a named list.

        Args:
            list_name: The list name, for example todo or shopping.
        """
        items = load_json("lists.json", {}).get(list_name.strip().lower(), [])
        return "\n".join(items) if items else "That list is empty."

    @function_tool
    async def remove_from_list(self, context: RunContext, list_name: str, item: str) -> str:
        """Remove items from a list that match some text. Use item "everything" to clear the list.

        Args:
            list_name: The list name.
            item: Text of the item to remove, or "everything".
        """
        lists = load_json("lists.json", {})
        key = list_name.strip().lower()
        old = lists.get(key, [])
        text = item.strip().lower()
        new = [] if text == "everything" else [i for i in old if text not in i.lower()]
        lists[key] = new
        save_json("lists.json", lists)
        return f"Removed {len(old) - len(new)} item(s)."

    @function_tool
    async def daily_briefing(self, context: RunContext, city: str = "") -> str:
        """Get today's weather, top news headlines, reminders, and to-do items for a briefing.

        Args:
            city: The user's city for the weather. Leave empty to use the saved default.
        """
        city = city.strip() or os.getenv("HOME_CITY", "")
        unit = os.getenv("WEATHER_UNITS", "fahrenheit")
        parts = []
        async with httpx.AsyncClient(timeout=15) as client:
            if city:
                try:
                    g = (await client.get("https://geocoding-api.open-meteo.com/v1/search",
                                          params={"name": city, "count": 1})).json()["results"][0]
                    w = (await client.get("https://api.open-meteo.com/v1/forecast", params={
                        "latitude": g["latitude"], "longitude": g["longitude"],
                        "current": "temperature_2m,weather_code",
                        "daily": "temperature_2m_max,temperature_2m_min,precipitation_probability_max",
                        "temperature_unit": unit, "timezone": "auto", "forecast_days": 1,
                    })).json()
                    cur, day = w["current"], w["daily"]
                    parts.append(
                        f"Weather in {g['name']}: {WEATHER_WORDS.get(cur['weather_code'], 'unsettled')}, "
                        f"now {round(cur['temperature_2m'])} degrees, high {round(day['temperature_2m_max'][0])}, "
                        f"low {round(day['temperature_2m_min'][0])}, chance of rain {day['precipitation_probability_max'][0]} percent.")
                except Exception as e:
                    parts.append(f"Weather lookup failed: {e}")
            else:
                parts.append("No city is set. Ask the user for their city and remember it.")
            try:
                rss = await client.get("http://feeds.bbci.co.uk/news/rss.xml")
                titles = [t.text for t in ET.fromstring(rss.text).findall(".//item/title")][:5]
                parts.append("Top headlines: " + "; ".join(titles))
            except Exception as e:
                parts.append(f"News lookup failed: {e}")
        reminders = load_json("reminders.json", [])
        if reminders:
            parts.append("Upcoming reminders: " + "; ".join(r["message"] for r in reminders))
        todo = load_json("lists.json", {}).get("todo", [])
        if todo:
            parts.append("To-do items: " + "; ".join(todo))
        return "\n".join(parts)

    @function_tool
    async def media_control(self, context: RunContext, action: str, amount: int = 1) -> str:
        """Control computer volume and media playback.

        Args:
            action: One of volume_up, volume_down, mute, play_pause, next, previous, stop.
            amount: For volume, how many ten percent steps (1 to 10).
        """
        if sys.platform != "win32":
            return "Media control only works on Windows."
        import ctypes
        action = action.strip().lower()
        if action in ("volume_up", "volume_down"):
            vk = 0xAF if action == "volume_up" else 0xAE
            presses = 5 * max(1, min(int(amount), 10))
        elif action in MEDIA_KEYS:
            vk, presses = MEDIA_KEYS[action], 1
        else:
            return "Unknown action."

        def _press():
            for _ in range(presses):
                ctypes.windll.user32.keybd_event(vk, 0, 0, 0)
                ctypes.windll.user32.keybd_event(vk, 0, 2, 0)
                time.sleep(0.02)

        await asyncio.to_thread(_press)
        return "Done."

    @function_tool
    async def research(self, context: RunContext, topic: str, open_report: bool = True) -> str:
        """Research a topic in depth from many web sources and save a written report.

        Args:
            topic: What to research, for example "best laptops under 1000 dollars".
            open_report: Open the finished report on the user's screen.
        """
        def _search(q):
            with DDGS() as d:
                return list(d.text(q, max_results=6))

        queries = [topic, f"{topic} guide", f"{topic} comparison review", f"{topic} latest"]
        batches = await asyncio.gather(*[asyncio.to_thread(_search, q) for q in queries], return_exceptions=True)
        seen, found = set(), []
        for batch in batches:
            if isinstance(batch, list):
                for r in batch:
                    url = r.get("href")
                    if url and url not in seen:
                        seen.add(url)
                        found.append(r)
        found = found[:8]
        if not found:
            return "I could not find anything on that topic."
        async with httpx.AsyncClient(timeout=10) as client:
            pages = await asyncio.gather(*[fetch_page(client, r["href"]) for r in found])
        sources = "\n\n".join(
            f"[{i}] {r.get('title', '')} ({r['href']})\n{text or r.get('body', '')}"
            for i, (r, text) in enumerate(zip(found, pages), 1))
        prompt = (
            f"Write a well-organized research report on: {topic}\n"
            "Use ONLY the numbered sources below and write everything in your own words. "
            "Format in Markdown: a title line starting with '# ', then '## Summary' (two or three sentences), "
            "'## Key findings' (bullets), any extra sections that help, '## Recommendation' if it makes sense, "
            "and '## Sources' listing each source number with its web address. "
            "Point out disagreements or uncertainty. Do not invent facts.\n\nSOURCES:\n" + sources)
        try:
            md = await gemini_text(prompt)
        except Exception as e:
            return f"Writing the report failed: {e}"
        title = next((l[2:].strip() for l in md.splitlines() if l.startswith("# ")), topic)
        path = save_document(title, md)
        if open_report and sys.platform == "win32":
            os.startfile(path)
        return f"Report saved as {path.name} in the documents folder. Summary: {section(md, 'Summary') or md[:600]}"

    @function_tool
    async def write_document(self, context: RunContext, title: str, instructions: str, open_it: bool = True) -> str:
        """Write a document such as an essay, letter, plan, or summary and save it as a file.

        Args:
            title: A short title for the document.
            instructions: Everything the document should cover, including tone and length.
            open_it: Open the finished document on the user's screen.
        """
        about = "\n".join(load_memories())
        system = "You are a skilled writer. Write clean, ready-to-use documents in Markdown."
        prompt = f"Write this document: {title}\n\nRequest: {instructions}\n"
        if about:
            prompt += f"\nFacts about the user that may help personalize it:\n{about}\n"
        try:
            md = await gemini_text(prompt, system)
        except Exception as e:
            return f"Writing failed: {e}"
        path = save_document(title, md)
        if open_it and sys.platform == "win32":
            os.startfile(path)
        return f"Saved {path.name} in the documents folder. It starts: {md[:300]}"

    @function_tool
    async def open_document(self, context: RunContext, name: str) -> str:
        """Open a saved report or document by part of its name.

        Args:
            name: A word from the document's name.
        """
        files = sorted(DOCS_DIR.glob("*"), key=lambda p: p.stat().st_mtime, reverse=True) if DOCS_DIR.exists() else []
        match = next((f for f in files if name.strip().lower() in f.name.lower()), None)
        if match is None:
            return "I could not find a saved document with that name."
        if sys.platform != "win32":
            return f"Found {match.name}, but opening only works on Windows."
        os.startfile(match)
        return f"Opened {match.name}."

    # ---------- computer control ----------
    def _pc_only(self) -> bool:
        if self._room is None:
            return False
        people = list(self._room.remote_participants.values())
        return bool(people) and all(p.attributes.get("source") == "local" for p in people)

    def _pc_only_reason(self) -> str:
        if self._room is None:
            return "Computer control is blocked: the agent has no room information. The agent file may be out of date."
        people = list(self._room.remote_participants.values())
        if any(p.attributes.get("source") == "phone" for p in people):
            return "Computer control is blocked because a phone is connected."
        return ("Computer control is blocked because the server did not mark this visitor as the PC. "
                "Tell the user to replace server.py with the newest version and restart everything.")

    def _guard(self):
        if not HAVE_GUI:
            return f"Computer control could not load. Reason: {GUI_ERROR or 'unknown'}. Tell the user to run the check file."
        if sys.platform != "win32":
            return "Computer control only works on Windows."
        if not self._pc_only():
            return self._pc_only_reason()
        if not self._armed:
            return "Computer control is off. Ask the user if you should turn it on."
        return None

    def _needs_confirm(self, key: str, confirmed: bool):
        if confirmed and self._pending == key:
            self._pending = None
            return None
        self._pending = key
        return (f"CONFIRMATION REQUIRED. Ask the user out loud: Should I {key}? "
                "Only after a clear yes, call this tool again with user_confirmed set to true.")

    @function_tool
    async def enable_control(self, context: RunContext) -> str:
        """Turn on computer control. Only call this when the user asks for it."""
        if not HAVE_GUI:
            return f"Computer control could not load. Reason: {GUI_ERROR or 'unknown'}. Tell the user to run the check file."
        if sys.platform != "win32":
            return "Computer control only works on Windows."
        if not self._pc_only():
            return self._pc_only_reason()
        self._armed = True
        return "Computer control is on. The user can say stop at any time."

    @function_tool
    async def stop_control(self, context: RunContext) -> str:
        """Turn off computer control immediately. Use this whenever the user says stop."""
        self._armed = False
        self._pending = None
        return "Computer control is off."

    @function_tool
    async def look_at_screen(self, context: RunContext, question: str = "Describe what is on the screen.") -> str:
        """Look at the user's screen and answer a question about it.

        Args:
            question: What to look for or read on the screen.
        """
        err = self._guard()
        if err:
            return err
        jpeg = await asyncio.to_thread(screenshot_jpeg)
        try:
            return await gemini_text(question + "\nAnswer briefly in plain sentences.", image_jpeg=jpeg)
        except Exception as e:
            return f"Could not read the screen: {e}"

    @function_tool
    async def click_on(self, context: RunContext, description: str, double: bool = False,
                       right: bool = False, user_confirmed: bool = False) -> str:
        """Find something on the screen by describing it, then click it.

        Args:
            description: What to click, for example "the blue Search button" or "the Chrome icon".
            double: Double click instead of a single click.
            right: Right click instead of a left click.
            user_confirmed: Set true only after the user said yes to a confirmation question.
        """
        err = self._guard()
        if err:
            return err
        if click_is_risky(description):
            need = self._needs_confirm(f"click {description}", user_confirmed)
            if need:
                return need
        jpeg = await asyncio.to_thread(screenshot_jpeg)
        prompt = (f'Find this on the screen: "{description}". Reply with ONLY JSON like '
                  '{"found": true, "x": 512, "y": 300} where x and y are the center of the item, '
                  'each a number from 0 to 1000 (0,0 is the top-left corner and 1000,1000 is the '
                  'bottom-right). If it is not visible reply {"found": false}.')
        try:
            point = parse_point(await gemini_text(prompt, image_jpeg=jpeg))
        except Exception as e:
            return f"Could not look at the screen: {e}"
        if point is None:
            return "I could not find that on the screen."
        sw, sh = pyautogui.size()
        x, y = int(point[0] * sw), int(point[1] * sh)

        def _click():
            pyautogui.moveTo(x, y, duration=0.3)
            if right:
                pyautogui.rightClick()
            elif double:
                pyautogui.doubleClick()
            else:
                pyautogui.click()
            time.sleep(0.4)

        await asyncio.to_thread(_click)
        return f"Clicked {description}."

    @function_tool
    async def type_text(self, context: RunContext, text: str) -> str:
        """Type text where the cursor currently is. Never use this for passwords.

        Args:
            text: The plain text to type.
        """
        err = self._guard()
        if err:
            return err
        if not text.isascii():
            return "I can only type plain letters, numbers, and common symbols."
        await asyncio.to_thread(pyautogui.write, text, 0.02)
        return "Typed it."

    @function_tool
    async def press_keys(self, context: RunContext, keys: str, user_confirmed: bool = False) -> str:
        """Press a key or shortcut such as enter, tab, ctrl+c, or ctrl+t.

        Args:
            keys: Key names joined with plus signs, for example "ctrl+shift+t".
            user_confirmed: Set true only after the user said yes to a confirmation question.
        """
        err = self._guard()
        if err:
            return err
        parts = [k.strip().lower() for k in keys.split("+") if k.strip()]
        if not parts or any(k not in pyautogui.KEYBOARD_KEYS for k in parts):
            return "I don't know one of those keys."
        if keys_are_risky(parts):
            need = self._needs_confirm(f"press {'+'.join(parts)}", user_confirmed)
            if need:
                return need
        await asyncio.to_thread(pyautogui.hotkey, *parts)
        return f"Pressed {'+'.join(parts)}."

    @function_tool
    async def scroll(self, context: RunContext, direction: str, amount: int = 5) -> str:
        """Scroll the page under the mouse.

        Args:
            direction: up or down.
            amount: How many notches, 1 to 20.
        """
        err = self._guard()
        if err:
            return err
        n = max(1, min(int(amount), 20)) * 120
        await asyncio.to_thread(pyautogui.scroll, n if direction.strip().lower() == "up" else -n)
        return "Scrolled."


server = AgentServer()


@server.rtc_session(agent_name="cortana")
async def entrypoint(ctx: JobContext):
    await ctx.connect()

    options = {"voice": "Kore"}
    if os.getenv("GEMINI_MODEL"):
        options["model"] = os.getenv("GEMINI_MODEL")

    session = AgentSession(llm=google.realtime.RealtimeModel(**options))

    await session.start(
        agent=Cortana(ctx.room),
        room=ctx.room,
        room_options=room_io.RoomOptions(video_input=True),
    )


if __name__ == "__main__":
    cli.run_app(server)
